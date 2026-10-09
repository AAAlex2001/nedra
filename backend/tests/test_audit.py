"""Тесты Блиц-аудита: заявка, торг по цене, группа, План аудита, договор и отчёт о документах."""

import asyncio
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import docx
import pytest
from pydantic import ValidationError

from app.models.billing import InvoiceStage
from app.models.expertise import (
    AuditTeamMember,
    ContractKind,
    CustomerType,
    Expertise,
    ExpertiseDocument,
    ExpertiseStatus,
    ServiceKind,
)
from app.models.user import UserRole
from app.schemas.audit import AuditInSchema, AuditOfferAnswerSchema, AuditPriceSchema
from app.schemas.audit_details import (
    AuditApplicantSchema,
    AuditBudgetSchema,
    AuditDetailsSchema,
    AuditFleetSchema,
    AuditObjectSchema,
    AuditParamsSchema,
    AuditPlanInSchema,
    AuditPlanSchema,
    AuditTimingSchema,
)
from app.services.audit.checklist import AUDIT_DOCUMENTS
from app.services.audit.invite import build_invite
from app.services.audit.report import build_documents_report
from app.services.audit.team import TeamMember
from app.services.audit.usecases.create_audit import CreateAuditUseCase
from app.services.audit.usecases.negotiate import (
    AnswerCounterUseCase,
    AnswerOfferUseCase,
    ProposePriceUseCase,
)
from app.services.audit.usecases.plan import (
    ApprovePlanUseCase,
    RequestPlanChangesUseCase,
    SendPlanUseCase,
)
from app.services.audit.usecases.team import SetAuditTeamUseCase
from app.services.audit.usecases.upload_documents import UploadAuditDocumentsUseCase
from app.services.contracts.document import CONTRACT, contract_number
from app.services.contracts.executors import NEDRA_AUDIT, executor_for
from app.services.contracts.kinds import allowed_kinds
from app.services.expertise.access import can_view
from app.services.expertise.exceptions import (
    ExpertiseAccessError,
    ExpertiseStateError,
    InvalidExpertiseError,
)
from app.services.expertise.presenter import executors_visible
from app.services.expertise.repo import executor_fits
from app.services.expertise.stages import mark_stage_paid
from app.services.expertise.usecases.accept_expertise import AcceptExpertiseUseCase
from app.services.expertise.usecases.create_expertise import build_company
from app.services.expertise.usecases.send_conclusion import SendConclusionUseCase
from test_expertise import (
    COMPANY,
    FITTING,
    FakeExpertiseRepository,
    FakeNotificationRepository,
    FakeProfileRepository,
    FakeStorage,
    FakeUpload,
    make_user,
    signed_text,
)

APPLICANT = AuditApplicantSchema(
    full_name="Иванов Иван Иванович",
    position="Главный инженер",
    organization="ООО «Ромашка»",
    inn="5405123456",
    phone="+79990000000",
    email="ivanov@romashka.ru",
)

OPO = AuditObjectSchema(
    reg_number="А59-12345-0001",
    name="Сеть газопотребления",
    hazard_class="III",
    address="г. Новосибирск, ул. Ленина, д. 1",
    industry="Газопотребление",
    hazard_signs=["substances", "pressure"],
)

BASIC = AuditParamsSchema(kind="basic", use_sto=False)
MONTH = AuditTimingSchema(kind="month")
FIXED = AuditBudgetSchema(mode="custom", negotiation="no")
OPEN = AuditBudgetSchema(mode="custom", negotiation="yes")
NO_BUDGET = AuditBudgetSchema(mode="none", negotiation="yes")


def make_details(**fields) -> AuditDetailsSchema:
    values = {
        "applicant": APPLICANT,
        "scope": "one",
        "objects": [OPO],
        "stages": ["documents"],
        "params": BASIC,
        "timing": MONTH,
        "budget": FIXED,
    }
    values.update(fields)

    return AuditDetailsSchema(**values)


def make_audit_order(details: AuditDetailsSchema | None = None, price: str | None = "150000") -> AuditInSchema:
    return AuditInSchema(
        details=details or make_details(),
        price=Decimal(price) if price else None,
        company=COMPANY,
    )


def make_audit(
    status: ExpertiseStatus,
    budget: AuditBudgetSchema = FIXED,
    price: Decimal | None = Decimal("150000"),
    stages: list[str] | None = None,
) -> Expertise:
    details = make_details(budget=budget, stages=stages or ["documents"])
    audit = Expertise(
        customer_id=1,
        service=ServiceKind.AUDIT,
        contract_kind=ContractKind.AUDIT,
        object_name="Сеть газопотребления (А59-12345-0001)",
        status=status,
        price=price,
        customer_type=CustomerType.LEGAL,
        audit_details=details.model_dump(mode="json"),
    )
    audit.id = 7
    return audit


def make_create_usecase(leads=None) -> tuple[CreateAuditUseCase, FakeNotificationRepository]:
    notifications = FakeNotificationRepository()
    usecase = CreateAuditUseCase(
        FakeExpertiseRepository(),
        FakeProfileRepository([], auditors=leads or [], leads=leads or []),
        notifications,
        FakeStorage(),
    )
    return usecase, notifications


PLAN = AuditPlanInSchema(
    documents_start=date(2026, 11, 2),
    documents_end=date(2026, 11, 13),
    meetings=True,
    opening_at=datetime(2026, 11, 2, 7, 0, tzinfo=timezone.utc),
    closing_at=datetime(2026, 11, 13, 7, 0, tzinfo=timezone.utc),
    interviewees="Главный инженер, специалист по промышленной безопасности",
)

TEAM = [TeamMember(user_id=20, full_name="Петров Пётр Петрович", areas=["Э1"], lead=True)]


def test_create_audit_keeps_items_and_notifies_leads() -> None:
    lead = make_user(20, UserRole.EXPERT)
    usecase, notifications = make_create_usecase([lead])

    files = [FakeUpload("устав.pdf"), FakeUpload("лицензия.pdf"), FakeUpload("лицензия-2.pdf")]
    created = asyncio.run(usecase.execute(make_user(1, UserRole.CUSTOMER), make_audit_order(), files, [1, 3, 3]))

    audit = created.expertise
    assert audit.service == ServiceKind.AUDIT
    assert audit.status == ExpertiseStatus.NEW
    assert audit.contract_kind == ContractKind.AUDIT
    assert audit.object_name == "Сеть газопотребления (А59-12345-0001)"
    assert audit.audit_details["params"]["kind"] == "basic"
    assert audit.audit_details["objects"][0]["hazard_signs"] == ["substances", "pressure"]
    assert [item.item_number for item in audit.documents] == [1, 3, 3]
    assert created.notified_experts == [lead]
    assert notifications.added[0].user_id == 20


def test_create_audit_allows_no_documents_but_checks_items() -> None:
    usecase, notifications = make_create_usecase()
    customer = make_user(1, UserRole.CUSTOMER)

    created = asyncio.run(usecase.execute(customer, make_audit_order(), [], []))
    assert created.expertise.documents == []

    with pytest.raises(InvalidExpertiseError):
        asyncio.run(usecase.execute(customer, make_audit_order(), [FakeUpload("a.pdf")], [48]))

    with pytest.raises(InvalidExpertiseError):
        asyncio.run(usecase.execute(customer, make_audit_order(), [FakeUpload("a.pdf")], []))


def test_create_audit_requires_proxy_and_sto_files() -> None:
    usecase, notifications = make_create_usecase()
    customer = make_user(1, UserRole.CUSTOMER)
    proxy = APPLICANT.model_copy(update={"by_proxy": True})
    sto = AuditParamsSchema(kind="interim", use_sto=True, sto_name="СТО 01-2025 «Аудит СУПБ»")

    with pytest.raises(InvalidExpertiseError):
        asyncio.run(usecase.execute(customer, make_audit_order(make_details(applicant=proxy)), [], []))

    with pytest.raises(InvalidExpertiseError):
        asyncio.run(usecase.execute(customer, make_audit_order(make_details(params=sto)), [], []))

    created = asyncio.run(
        usecase.execute(
            customer,
            make_audit_order(make_details(applicant=proxy, params=sto)),
            [],
            [],
            power_of_attorney=FakeUpload("доверенность.pdf"),
            sto_file=FakeUpload("сто.pdf"),
        )
    )
    assert [item.kind for item in created.expertise.documents] == ["power_of_attorney", "sto"]


def test_audit_of_all_objects_keeps_certificate() -> None:
    usecase, notifications = make_create_usecase()
    fleet = AuditFleetSchema(count=12, profile="Угольная промышленность", multi_region=True)
    details = make_details(scope="all", objects=[], fleet=fleet)

    created = asyncio.run(
        usecase.execute(
            make_user(1, UserRole.CUSTOMER),
            make_audit_order(details),
            [FakeUpload("устав.pdf")],
            [1],
            opo_certificate=FakeUpload("свидетельство.pdf"),
        )
    )

    audit = created.expertise
    assert audit.object_name == "Все ОПО ООО «Ромашка» (12)"
    assert [item.kind for item in audit.documents] == ["audit_item", "opo_certificate"]


def test_consultation_goes_to_manager_first() -> None:
    lead = make_user(20, UserRole.EXPERT)
    usecase, notifications = make_create_usecase([lead])
    details = make_details(params=AuditParamsSchema(kind="consultation"))

    created = asyncio.run(usecase.execute(make_user(1, UserRole.CUSTOMER), make_audit_order(details), [], []))

    assert created.expertise.status == ExpertiseStatus.CONSULTATION
    assert created.notified_experts == []
    assert notifications.added == []


def test_audit_details_validation() -> None:
    with pytest.raises(ValidationError):
        make_details(scope="one", objects=[])

    with pytest.raises(ValidationError):
        make_details(scope="selected", objects=[])

    with pytest.raises(ValidationError):
        make_details(scope="all", objects=[])

    with pytest.raises(ValidationError):
        make_details(stages=[])

    with pytest.raises(ValidationError):
        AuditParamsSchema(kind="basic")

    with pytest.raises(ValidationError):
        AuditParamsSchema(kind="interim", use_sto=True, sto_name=" ")

    with pytest.raises(ValidationError):
        AuditParamsSchema(kind="selective", elements=[])

    with pytest.raises(ValidationError):
        AuditTimingSchema(kind="custom", start=date(2026, 12, 1))

    with pytest.raises(ValidationError):
        AuditTimingSchema(kind="custom", start=date(2026, 12, 10), end=date(2026, 12, 1))

    with pytest.raises(ValidationError):
        AuditObjectSchema(**(OPO.model_dump() | {"hazard_signs": []}))


def test_audit_price_and_customer_rules() -> None:
    with pytest.raises(ValidationError):
        make_audit_order(price="99999")

    with pytest.raises(ValidationError):
        make_audit_order(price=None)

    unset = make_audit_order(make_details(budget=NO_BUDGET), price=None)
    assert unset.price is None

    with pytest.raises(ValidationError):
        make_audit_order(make_details(budget=NO_BUDGET), price="500000")

    with pytest.raises(ValidationError):
        AuditInSchema(details=make_details(), price=Decimal("150000"), company=COMPANY, customer_type=CustomerType.INDIVIDUAL)

    entrepreneur = AuditInSchema(
        details=make_details(), price=Decimal("150000"), company=COMPANY, customer_type=CustomerType.ENTREPRENEUR
    )
    assert entrepreneur.customer_type == CustomerType.ENTREPRENEUR

    with pytest.raises(ValidationError):
        AuditPriceSchema(price=Decimal("90000"))

    with pytest.raises(ValidationError):
        AuditOfferAnswerSchema(answer="counter")


def test_audit_is_visible_only_to_leads() -> None:
    audit = make_audit(ExpertiseStatus.NEW)
    expert = make_user(10, UserRole.EXPERT)

    assert executor_fits(audit, [], True)
    assert not executor_fits(audit, FITTING, False)
    assert can_view(audit, expert, [], True)
    assert not can_view(audit, expert, FITTING, False)

    audit.expert_id = 20
    audit.team = [AuditTeamMember(user_id=10)]
    assert can_view(audit, expert, [], False)


def test_accept_audit_requires_lead_and_hides_name() -> None:
    auditor = make_user(30, UserRole.EXPERT)
    lead = make_user(20, UserRole.EXPERT)
    lead.full_name = "Петров Пётр Петрович"
    profiles = FakeProfileRepository([], FITTING, auditors=[auditor, lead], leads=[lead])
    notifications = FakeNotificationRepository()
    usecase = AcceptExpertiseUseCase(FakeExpertiseRepository(), profiles, notifications)

    with pytest.raises(ExpertiseAccessError):
        asyncio.run(usecase.execute(auditor, make_audit(ExpertiseStatus.NEW)))

    accepted = asyncio.run(usecase.execute(lead, make_audit(ExpertiseStatus.NEW)))
    assert accepted.status == ExpertiseStatus.EXPERT_READY
    assert "Петров" not in notifications.added[0].text
    assert "Аудитор НПИ «Недра»" in notifications.added[0].text


def test_customer_sees_auditors_only_after_plan() -> None:
    audit = make_audit(ExpertiseStatus.CONTRACT)
    customer = make_user(1, UserRole.CUSTOMER)
    lead = make_user(20, UserRole.EXPERT)

    assert not executors_visible(audit, customer)
    assert executors_visible(audit, lead)

    audit.plan_sent_at = datetime(2026, 10, 20, tzinfo=timezone.utc)
    assert executors_visible(audit, customer)


def test_fixed_budget_cannot_be_negotiated() -> None:
    lead = make_user(20, UserRole.EXPERT)
    usecase = ProposePriceUseCase(
        FakeExpertiseRepository(), FakeProfileRepository([], leads=[lead]), FakeNotificationRepository()
    )

    with pytest.raises(ExpertiseStateError):
        asyncio.run(usecase.execute(lead, make_audit(ExpertiseStatus.NEW, FIXED), Decimal("200000")))


def test_negotiation_round_ends_in_agreement() -> None:
    lead = make_user(20, UserRole.EXPERT)
    customer = make_user(1, UserRole.CUSTOMER)
    expertises = FakeExpertiseRepository()
    notifications = FakeNotificationRepository()
    audit = make_audit(ExpertiseStatus.NEW, NO_BUDGET, price=None)

    propose = ProposePriceUseCase(expertises, FakeProfileRepository([], leads=[lead]), notifications)
    asyncio.run(propose.execute(lead, audit, Decimal("300000")))
    assert audit.status == ExpertiseStatus.OFFER
    assert audit.expert_id == 20
    assert notifications.added[-1].user_id == 1

    answer = AnswerOfferUseCase(expertises, notifications)
    with pytest.raises(InvalidExpertiseError):
        asyncio.run(answer.execute(customer, audit, "counter", Decimal("350000")))

    asyncio.run(answer.execute(customer, audit, "counter", Decimal("250000")))
    assert audit.status == ExpertiseStatus.COUNTER
    assert audit.counter_price == Decimal("250000")

    counter = AnswerCounterUseCase(expertises, notifications)
    asyncio.run(counter.execute(lead, audit, True))
    assert audit.status == ExpertiseStatus.EXPERT_READY
    assert audit.price == Decimal("250000")
    assert audit.offer_price is None


def test_declined_offer_returns_to_pool() -> None:
    lead = make_user(20, UserRole.EXPERT)
    customer = make_user(1, UserRole.CUSTOMER)
    expertises = FakeExpertiseRepository()
    notifications = FakeNotificationRepository()
    audit = make_audit(ExpertiseStatus.NEW, OPEN)

    propose = ProposePriceUseCase(expertises, FakeProfileRepository([], leads=[lead]), notifications)
    asyncio.run(propose.execute(lead, audit, Decimal("200000")))

    asyncio.run(AnswerOfferUseCase(expertises, notifications).execute(customer, audit, "decline"))
    assert audit.status == ExpertiseStatus.NEW
    assert audit.expert_id is None
    assert audit.offer_price is None

    asyncio.run(propose.execute(lead, audit, Decimal("180000")))
    asyncio.run(AnswerOfferUseCase(expertises, notifications).execute(customer, audit, "counter", Decimal("150000")))
    asyncio.run(AnswerCounterUseCase(expertises, notifications).execute(lead, audit, False))
    assert audit.status == ExpertiseStatus.NEW
    assert audit.expert_id is None


def test_advance_moves_audit_to_plan() -> None:
    audit = make_audit(ExpertiseStatus.CONTRACT)

    text = mark_stage_paid(audit, InvoiceStage.ADVANCE)

    assert audit.status == ExpertiseStatus.PLAN
    assert "План аудита" in text


def test_lead_forms_team_of_auditors() -> None:
    lead = make_user(20, UserRole.EXPERT)
    auditor = make_user(30, UserRole.EXPERT)
    notifications = FakeNotificationRepository()
    profiles = FakeProfileRepository([], auditors=[lead, auditor], leads=[lead])
    usecase = SetAuditTeamUseCase(FakeExpertiseRepository(), profiles, notifications)
    audit = make_audit(ExpertiseStatus.PLAN)
    audit.expert_id = 20

    asyncio.run(usecase.execute(lead, audit, [30, 20, 30]))
    assert [member.user_id for member in audit.team] == [30]
    assert notifications.added[0].user_id == 30

    with pytest.raises(InvalidExpertiseError):
        asyncio.run(usecase.execute(lead, audit, [99]))

    audit.status = ExpertiseStatus.PLAN_REVIEW
    with pytest.raises(ExpertiseStateError):
        asyncio.run(usecase.execute(lead, audit, [30]))


def test_plan_is_sent_approved_and_opens_documents() -> None:
    lead = make_user(20, UserRole.EXPERT)
    customer = make_user(1, UserRole.CUSTOMER)
    expertises = FakeExpertiseRepository()
    notifications = FakeNotificationRepository()
    storage = FakeStorage()
    audit = make_audit(ExpertiseStatus.PLAN)
    audit.expert_id = 20
    audit.company = build_company(COMPANY)

    upload = UploadAuditDocumentsUseCase(expertises, notifications, storage)
    with pytest.raises(ExpertiseStateError):
        asyncio.run(upload.execute(customer, audit, [FakeUpload("устав.pdf")], [1]))

    asyncio.run(SendPlanUseCase(expertises, notifications, storage).execute(lead, audit, PLAN, TEAM))
    assert audit.status == ExpertiseStatus.PLAN_REVIEW
    assert audit.audit_plan["version"] == 1
    assert audit.documents[-1].kind == "audit_plan"

    program = docx.Document(BytesIO(storage.written[-1]))
    cells = [cell.text for table in program.tables for row in table.rows for cell in row.cells]
    assert "ООО «Ромашка»" in "\n".join(cells)
    assert "Петров Пётр Петрович, области аттестации: Э1" in cells
    assert "А59-12345-0001" in cells
    assert not any("{{" in text for text in cells)

    asyncio.run(RequestPlanChangesUseCase(expertises, notifications).execute(customer, audit, "Перенесите на неделю"))
    assert audit.status == ExpertiseStatus.PLAN
    assert audit.plan_comment == "Перенесите на неделю"

    asyncio.run(SendPlanUseCase(expertises, notifications, storage).execute(lead, audit, PLAN, TEAM))
    assert audit.audit_plan["version"] == 2
    assert audit.plan_comment is None

    asyncio.run(ApprovePlanUseCase(expertises, notifications).execute(customer, audit))
    assert audit.status == ExpertiseStatus.IN_PROGRESS
    assert audit.plan_approved_at is not None

    asyncio.run(upload.execute(customer, audit, [FakeUpload("устав.pdf")], [1]))
    assert audit.documents[-1].kind == "audit_item"


def test_plan_requires_onsite_dates_when_chosen() -> None:
    lead = make_user(20, UserRole.EXPERT)
    audit = make_audit(ExpertiseStatus.PLAN, stages=["documents", "onsite"])
    audit.expert_id = 20
    audit.company = build_company(COMPANY)
    usecase = SendPlanUseCase(FakeExpertiseRepository(), FakeNotificationRepository(), FakeStorage())

    with pytest.raises(InvalidExpertiseError):
        asyncio.run(usecase.execute(lead, audit, PLAN, TEAM))


def test_invite_is_built_for_opening_meeting() -> None:
    audit = make_audit(ExpertiseStatus.IN_PROGRESS)
    plan = AuditPlanSchema(**PLAN.model_dump(), version=1)

    invite = build_invite(audit, plan, datetime(2026, 10, 20, tzinfo=timezone.utc)).decode()

    assert "BEGIN:VEVENT" in invite
    assert "DTSTART:20261102T070000Z" in invite
    assert "DTEND:20261102T080000Z" in invite

    silent = AuditPlanSchema(**(PLAN.model_dump() | {"meetings": False}), version=1)
    assert build_invite(audit, silent, datetime(2026, 10, 20, tzinfo=timezone.utc)) is None


def test_audit_contract_uses_own_template_and_number() -> None:
    audit = make_audit(ExpertiseStatus.EXPERT_READY)
    audit.expert_id = 20
    audit.deadline = "week"
    audit.company = build_company(COMPANY)

    signed_at = datetime(2026, 10, 6, tzinfo=timezone.utc)

    assert executor_for(ContractKind.AUDIT) == NEDRA_AUDIT
    assert contract_number(audit, signed_at) == "БА-2026-0007"
    assert ContractKind.AUDIT not in allowed_kinds(None)

    text = signed_text(CONTRACT, audit)
    assert "аудита системы управления промышленной безопасностью" in text
    assert "Сеть газопотребления (А59-12345-0001)" in text


def test_send_audit_report_without_result() -> None:
    audit = make_audit(ExpertiseStatus.PAID)
    audit.expert_id = 20
    usecase = SendConclusionUseCase(FakeExpertiseRepository(), FakeNotificationRepository(), FakeStorage())

    sent = asyncio.run(usecase.execute(make_user(20, UserRole.EXPERT), audit, None, [FakeUpload("отчёт.pdf")]))

    assert sent.status == ExpertiseStatus.SENT
    assert sent.result is None


def test_documents_report_counts_provided_items() -> None:
    audit = make_audit(ExpertiseStatus.NEW)
    audit.company = build_company(COMPANY)
    audit.documents = [
        ExpertiseDocument(kind="audit_item", item_number=1, uploaded_by=1, file_path="x", original_name="устав.pdf", size=1, content_type="application/pdf"),
        ExpertiseDocument(kind="audit_item", item_number=4, uploaded_by=1, file_path="x", original_name="свидетельство.pdf", size=1, content_type="application/pdf"),
    ]

    content = build_documents_report(audit, datetime(2026, 10, 6, 12, 0, tzinfo=timezone.utc))
    document = docx.Document(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert f"по 2 из {len(AUDIT_DOCUMENTS)} пунктов" in text
    table = document.tables[0]
    assert len(table.rows) == len(AUDIT_DOCUMENTS) + 1
    assert table.rows[1].cells[2].text == "Представлен"
    assert table.rows[1].cells[3].text == "устав.pdf"
    assert table.rows[2].cells[2].text == "Не представлен"
