"""Отчёт о представленных документах для аудитора: какие пункты перечня закрыты файлами, а какие нет."""

from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

from app.models.expertise import Expertise, ExpertiseDocument
from app.services.audit.checklist import AUDIT_DOCUMENTS

PROVIDED = "Представлен"
MISSING = "Не представлен"


def customer_title(audit: Expertise) -> str:
    """Заказчик для шапки отчёта: организация или физлицо."""

    if audit.company is not None:
        return f"{audit.company.name}, ИНН {audit.company.inn}"

    if audit.individual is not None:
        return audit.individual.full_name

    return "—"


def files_by_item(documents: list[ExpertiseDocument]) -> dict[int, list[str]]:
    """Имена файлов по номерам пунктов перечня."""

    grouped: dict[int, list[str]] = {}

    for document in documents:
        if document.kind != "audit_item" or document.item_number is None:
            continue
        if document.item_number not in grouped:
            grouped[document.item_number] = []
        grouped[document.item_number].append(document.original_name)

    return grouped


def build_documents_report(audit: Expertise, generated_at: datetime) -> bytes:
    """Собрать Word-файл: шапка, итог по комплектности и таблица по всем пунктам перечня."""

    grouped = files_by_item(audit.documents)
    provided = len(grouped)
    total = len(AUDIT_DOCUMENTS)

    document = Document()

    section = document.sections[0]
    portrait_width = section.page_width
    portrait_height = section.page_height
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = portrait_height
    section.page_height = portrait_width
    section.left_margin = Cm(2)
    section.right_margin = Cm(1.5)

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    document.add_heading("Отчёт о представленных документах", level=1)
    document.add_paragraph(
        "для проведения оценки соблюдения обязательных требований "
        "(аудит системы управления промышленной безопасностью)"
    )

    document.add_paragraph(f"Заявка № {audit.id}")
    document.add_paragraph(f"Заказчик: {customer_title(audit)}")
    document.add_paragraph(f"Объект аудита: {audit.object_name or '—'}")
    document.add_paragraph(f"Дата формирования: {generated_at:%d.%m.%Y %H:%M}")

    summary = document.add_paragraph()
    summary.add_run(
        f"Представлены документы по {provided} из {total} пунктов перечня, "
        f"не представлены — по {total - provided}."
    ).bold = True

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    header = table.rows[0].cells
    header[0].text = "№"
    header[1].text = "Документ по перечню"
    header[2].text = "Статус"
    header[3].text = "Файлы"

    for number, title in enumerate(AUDIT_DOCUMENTS, start=1):
        names = grouped.get(number, [])
        row = table.add_row().cells
        row[0].text = str(number)
        row[1].text = title
        row[2].text = PROVIDED if names else MISSING
        row[3].text = "\n".join(names)

    widths = (Cm(1), Cm(13), Cm(3), Cm(8))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width

    revisions = [item for item in audit.documents if item.kind == "revision"]
    if revisions:
        document.add_heading("Документы, присланные после замечаний", level=2)
        for item in revisions:
            document.add_paragraph(item.original_name, style="List Bullet")

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def report_filename(audit: Expertise) -> str:
    """Имя файла отчёта для скачивания."""

    return f"Отчёт о представленных документах — заявка {audit.id}.docx"
