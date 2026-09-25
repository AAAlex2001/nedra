"""Заполнение договора и соглашения о конфиденциальности реквизитами заявки.

Шаблоны — документы юристов в Word, где вместо пропусков стоят метки вида
{{customer_inn}}. Каждая метка лежит целиком внутри одного фрагмента текста,
поэтому замена не ломает оформление документа. Оба документа подписываются
одновременно и ссылаются на один номер договора.
"""

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as WordDocument
from docx.text.paragraph import Paragraph

from app.models.expertise import ContractKind, Expertise
from app.models.user import User
from app.services.contracts.executors import executor_for
from app.services.contracts.kinds import SUBJECTS
from app.services.documents.layout import (
    amount_in_words,
    format_amount,
    number_to_words,
    plural,
    vat_included,
)

TEMPLATES = Path(__file__).parent / "templates"

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

CONTRACT = "contract"
NDA = "nda"
SIGNED_TITLES = {CONTRACT: "Договор", NDA: "Соглашение о конфиденциальности"}

MONTHS = (
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
)

WORKING_DAYS = ("рабочий день", "рабочих дня", "рабочих дней")

DEADLINE_DAYS = {"today": 1, "three_days": 3, "week": 5, "any": 5}
DEFAULT_DAYS = 5


def contract_problem(expertise: Expertise) -> str | None:
    """Чего не хватает, чтобы составить договор. None — всё на месте."""

    if expertise.company is None:
        return "В заявке нет реквизитов заказчика"

    if expertise.contract_kind is None:
        return "Вид договора ещё не определён"

    if not expertise.object_name:
        return "В заявке нет наименования документации"

    if expertise.price is None:
        return "Стоимость экспертизы не задана"

    return None


def contract_number(expertise: Expertise, signed_at: datetime) -> str:
    """Номер договора: префикс сервиса, год подписания и номер заявки."""

    return f"БЭ-{signed_at.year}-{expertise.id:04d}"


def signed_filename(kind: str, expertise: Expertise, signed_at: datetime) -> str:
    """Имя файла для скачивания: «Договор БЭ-2026-0001.docx»."""

    return f"{SIGNED_TITLES[kind]} {contract_number(expertise, signed_at)}.docx"


def format_date(moment: datetime) -> str:
    """Дата как в договоре: «25» сентября 2026 г."""

    return f"«{moment.day:02d}» {MONTHS[moment.month - 1]} {moment.year} г."


def initials(full_name: str) -> str:
    """Иванов Иван Иванович → И.И. Иванов."""

    surname, *names = full_name.split()
    letters = "".join(f"{name[0]}." for name in names)

    return f"{letters} {surname}" if letters else surname


def working_days(deadline: str | None) -> str:
    """Срок из пожелания заказчика: «3 (три) рабочих дня»."""

    days = DEADLINE_DAYS.get(deadline or "", DEFAULT_DAYS)

    return f"{days} ({number_to_words(days)}) {plural(days, WORKING_DAYS)}"


def money_text(amount: Decimal) -> str:
    """Сумма цифрами и прописью: «20 000,00 (Двадцать тысяч рублей 00 копеек)»."""

    return f"{format_amount(amount)} ({amount_in_words(amount)})"


def vat_clause(amount: Decimal, rate: int) -> str:
    """Фраза про НДС в пункте о стоимости."""

    if rate <= 0:
        return "НДС не облагается"

    return f"в т.ч. НДС {rate}% – {money_text(vat_included(amount, rate))}"


def document_values(
    expertise: Expertise, customer: User, vat_rate: int, signed_at: datetime
) -> dict[str, str]:
    """Значения меток шаблонов. Перед вызовом заявку проверяет contract_problem.

    Телефон и почту для связи берём из аккаунта заказчика, который подписывает.
    """

    company = expertise.company
    if company is None:
        raise ValueError("В заявке нет реквизитов заказчика")

    price = expertise.price if expertise.price is not None else Decimal("0")
    kind = ContractKind(expertise.contract_kind)
    date = format_date(signed_at)

    return {
        "number": contract_number(expertise, signed_at),
        "date": date,
        "date_caps": date.upper(),
        "subject": SUBJECTS[kind],
        "object_name": expertise.object_name or "",
        "days": working_days(expertise.deadline),
        "price": money_text(price),
        "vat": vat_clause(price, vat_rate),
        "customer_full_name": company.full_name,
        "customer_name": company.name,
        "customer_inn": company.inn,
        "customer_kpp": company.kpp or "—",
        "customer_ogrn": company.ogrn,
        "customer_address": company.address,
        "customer_bank": company.bank,
        "customer_bic": company.bic,
        "customer_account": company.account,
        "customer_corr_account": company.corr_account,
        "customer_phone": customer.phone,
        "customer_email": customer.email,
        "signer_position": company.signer_position,
        "signer_genitive": company.signer_genitive,
        "signer_basis": company.signer_basis,
        "signer_initials": initials(company.signer_name),
    }


def all_paragraphs(document: WordDocument) -> list[Paragraph]:
    """Абзацы документа вместе с абзацами в ячейках таблиц."""

    cells = [
        paragraph
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    ]

    return [*document.paragraphs, *cells]


def fill(document: WordDocument, values: dict[str, str]) -> None:
    """Подставить значения вместо меток."""

    for paragraph in all_paragraphs(document):
        for run in paragraph.runs:
            if "{{" not in run.text:
                continue

            text = run.text
            for key, value in values.items():
                text = text.replace(f"{{{{{key}}}}}", value)
            run.text = text


def build_signed_document(
    kind: str, expertise: Expertise, customer: User, vat_rate: int, signed_at: datetime
) -> bytes:
    """Собрать договор или соглашение о конфиденциальности по шаблону исполнителя."""

    executor = executor_for(expertise.contract_kind)
    template = executor.contract_template if kind == CONTRACT else executor.nda_template
    document = Document(str(TEMPLATES / template))

    fill(document, document_values(expertise, customer, vat_rate, signed_at))

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()
